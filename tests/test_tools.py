"""
Tests de tools/ con un exporte inventado (sin datos reales).
"""
from datetime import date

import openpyxl
import pytest

from tools.aplicar_correcciones import aplicar, CorreccionInvalida
from tools.auditoria_excel import auditar

HEADERS_PERMISOS = ["id", "fecha", "legajo", "nombre", "hora_salida", "hora_entrada", "sin_retorno",
                    "motivo", "compensa", "minutos_reales", "horas_redondeadas", "registrado_por",
                    "planta", "timestamp"]


@pytest.fixture
def exporte(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("config_horarios")
    ws.append(["id", "planta", "desde", "hasta", "hora_fin", "descripcion", "activo",
               "excluidos_legajos", "creado_por", "timestamp"])
    ws.append(["H1", "Fábrica", "2026-09-01", "2026-12-31", "11:00", "Reducción", "SI", "2002", "RRHH", ""])
    ws = wb.create_sheet("padron")
    ws.append(["legajo", "nombre", "sector", "centro_costo", "planta", "activo", "clasificacion"])
    ws.append(["2001", "PEREZ, ANA", "COSTURA SAN JUAN", "", "Fábrica", "SI", "HOURLY DIRECT"])
    ws.append([2002, "GOMEZ, LUIS", "CALIDAD", "", "Fábrica", "SI", "EXEMPT"])
    ws.append(["2003", "DIAZ  EVA", "COSTURA SAN JUAN", "", "NO", "SI", "HOURLY DIRECT"])
    ws = wb.create_sheet("permisos")
    ws.append(HEADERS_PERMISOS)
    # S/R cargado contra las 15:00 con horario 6 a 11 vigente → debe marcarse
    ws.append(["P20260917131851", "2026-09-17", "2001", "PEREZ, ANA", "10:05", "S/R", "SI",
               "Banco / Cajero", "SI", 295, 5, "Guardia", "Fábrica", ""])
    # Excluido del horario especial: 10:00 → 15:00 está bien
    ws.append(["P20260918100000", "2026-09-18", 2002, "GOMEZ, LUIS", "10:00", "S/R", "SI",
               "Banco / Cajero", "SI", 300, 5, "Guardia", "Fábrica", ""])
    # Motivo de texto libre ("Otro"): compensa por política, no se marca
    ws.append(["P20260919100000", "2026-09-19", "2001", "PEREZ, ANA", "08:00", "09:00", "NO",
               "Registro Automotor", "SI", 60, 1, "Guardia", "Fábrica", ""])
    # Motivo oficial que no compensa en SJ → se marca
    ws.append(["P20260920100000", "2026-09-20", "2001", "PEREZ, ANA", "08:00", "09:00", "NO",
               "Trámite personal", "SI", 60, 1, "Guardia", "Fábrica", ""])
    ws = wb.create_sheet("compensaciones")
    ws.append(["id", "fecha_compensacion", "legajo", "nombre", "horas_compensadas", "observacion",
               "registrado_por", "planta", "timestamp"])
    ruta = tmp_path / "exporte.xlsx"
    wb.save(ruta)
    return ruta


def test_auditoria_usa_horario_vigente_y_excluidos(exporte):
    r = auditar(str(exporte), hoy=date(2026, 10, 19))
    sr = r[r["problema"].str.startswith("S/R")]
    assert list(sr["id"]) == ["P20260917131851"]
    assert "55min/1h" in sr.iloc[0]["detalle"]


def test_auditoria_politica_solo_motivos_oficiales(exporte):
    r = auditar(str(exporte), hoy=date(2026, 10, 19))
    assert list(r[r["problema"].str.contains("política")]["id"]) == ["P20260920100000"]


def test_auditoria_padron(exporte):
    r = auditar(str(exporte), hoy=date(2026, 10, 19))
    padron = r[r["hoja"] == "padron"]
    assert set(padron["problema"]) == {"planta rara", "espacios de más en el nombre"}
    assert set(padron["id"]) == {"2003"}


def _corr(**kw):
    base = {"hoja": "permisos", "id": "P20260917131851", "columna": "minutos_reales",
            "valor_anterior": "295", "valor_nuevo": "55", "motivo": "S/R"}
    return {**base, **kw}


def test_aplicar_por_nombre_de_columna_y_sin_tocar_original(exporte, tmp_path):
    salida = tmp_path / "corregido.xlsx"
    reg = aplicar(str(exporte), [
        _corr(),
        _corr(columna="registrado_por", valor_anterior="*", valor_nuevo="{actual} | corregido"),
        _corr(hoja="padron", id="2002", columna="activo", valor_anterior="SI", valor_nuevo="NO"),
    ], str(salida), hoy=date(2026, 10, 19))
    assert [x[1] for x in reg] == ["J2", "L2", "F3"]
    wb = openpyxl.load_workbook(salida)
    assert wb.sheetnames[0] == "CAMBIOS_19-10-2026"
    assert wb["permisos"]["J2"].value == 55.0
    assert wb["permisos"]["L2"].value == "Guardia | corregido"
    assert wb["padron"]["F3"].value == "NO" and wb["padron"]["E3"].value == "Fábrica"
    assert openpyxl.load_workbook(exporte)["permisos"]["J2"].value == 295


def test_aplicar_valor_anterior_distinto_no_guarda_nada(exporte, tmp_path):
    salida = tmp_path / "corregido.xlsx"
    with pytest.raises(CorreccionInvalida, match="se esperaba '300'"):
        aplicar(str(exporte), [_corr(columna="horas_redondeadas", valor_anterior="5", valor_nuevo="1"),
                               _corr(valor_anterior="300")], str(salida))
    assert not salida.exists()


def test_aplicar_columna_o_id_inexistente(exporte, tmp_path):
    with pytest.raises(CorreccionInvalida, match="no tiene la columna"):
        aplicar(str(exporte), [_corr(columna="horas")], str(tmp_path / "x.xlsx"))
    with pytest.raises(CorreccionInvalida, match="no se encontró"):
        aplicar(str(exporte), [_corr(id="P0")], str(tmp_path / "x.xlsx"))
