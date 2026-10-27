"""
tools/aplicar_correcciones.py
─────────────────────────────────────────────────────────────────
Aplica una lista de correcciones a una COPIA de un exporte del Google
Sheet (.xlsx), para después pegar las celdas cambiadas a mano en el
Sheet. Nunca toca el Sheet real ni el archivo original.

Uso (desde la raíz del repo):
    python -X utf8 tools/aplicar_correcciones.py export.xlsx correcciones.csv
    python -X utf8 tools/aplicar_correcciones.py export.xlsx correcciones.csv --salida corregido.xlsx

correcciones.csv (UTF-8, con encabezado; queda fuera del repo por .gitignore):

    hoja,id,columna,valor_anterior,valor_nuevo,motivo
    permisos,P20260917131851,minutos_reales,295,55,S/R recalculado con horario 6 a 11
    permisos,P20260917131851,horas_redondeadas,5,1,S/R recalculado con horario 6 a 11
    permisos,P20260917131851,registrado_por,*,{actual} | Corregido por ... — 19/10/2026,S/R recalculado
    padron,9348,activo,SI,NO,Baja

Reglas (las mismas que respeta la app, ver CLAUDE.md §9):
- La fila se busca por el valor de la columna A (id del registro; legajo en el padrón).
- La columna se busca por NOMBRE de encabezado, nunca por posición.
- valor_anterior tiene que coincidir con lo que hay hoy en la celda (o ser *
  para no verificar). Si una sola corrección no coincide, no se guarda nada:
  así no se pisa un dato que alguien cambió después del exporte.
- {actual} en valor_nuevo se reemplaza por el contenido actual de la celda
  (sirve para agregar texto de auditoría en registrado_por sin perder el guardia).
- (vacío) en valor_anterior / valor_nuevo significa celda vacía.

Resultado: una copia con las celdas cambiadas pintadas de amarillo y una
hoja CAMBIOS_<fecha> al principio con hoja, celda real, id, valor anterior,
valor nuevo y motivo: es la lista para pegar en el Sheet.
"""
import argparse
import csv
import os
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Font, PatternFill

AMARILLO = PatternFill("solid", fgColor="FFF59D")
VACIO = "(vacío)"
COLUMNAS_NUMERICAS = {"minutos_reales", "horas_redondeadas", "horas_compensadas"}


class CorreccionInvalida(Exception):
    pass


def _texto(valor) -> str:
    """Representación comparable de una celda: 2177, 2177.0 y '2177' son iguales."""
    if valor is None or valor == "":
        return VACIO
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d") if valor.time() == datetime.min.time() else valor.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(valor, date):
        return valor.strftime("%Y-%m-%d")
    if hasattr(valor, "hour"):  # datetime.time
        return valor.strftime("%H:%M")
    texto = str(valor).strip()
    try:
        numero = float(texto.replace(",", "."))
        return f"{numero:g}"
    except ValueError:
        return texto


def _valor_a_escribir(columna: str, texto: str):
    if texto == VACIO:
        return None
    if columna in COLUMNAS_NUMERICAS:
        try:
            return float(texto.replace(",", "."))
        except ValueError:
            raise CorreccionInvalida(f"'{columna}' es numérica y recibió '{texto}'")
    return texto


def leer_correcciones(ruta_csv: str) -> list[dict]:
    with open(ruta_csv, encoding="utf-8-sig", newline="") as fh:
        filas = list(csv.DictReader(fh))
    requeridas = {"hoja", "id", "columna", "valor_anterior", "valor_nuevo", "motivo"}
    if filas and not requeridas.issubset(filas[0].keys()):
        raise CorreccionInvalida(f"El CSV necesita las columnas: {', '.join(sorted(requeridas))}")
    return [{k: (v or "").strip() for k, v in f.items()} for f in filas]


def aplicar(ruta_excel: str, correcciones: list[dict], ruta_salida: str, hoy: date | None = None) -> list[tuple]:
    hoy = hoy or date.today()
    wb = openpyxl.load_workbook(ruta_excel)
    cambios, errores = [], []

    for n, c in enumerate(correcciones, start=2):  # n = línea del CSV
        try:
            if c["hoja"] not in wb.sheetnames:
                raise CorreccionInvalida(f"no existe la hoja '{c['hoja']}'")
            ws = wb[c["hoja"]]
            headers = [str(x.value).strip() if x.value is not None else "" for x in ws[1]]
            if c["columna"] not in headers:
                raise CorreccionInvalida(f"la hoja '{c['hoja']}' no tiene la columna '{c['columna']}'")
            col = headers.index(c["columna"]) + 1

            filas = [r for r in range(2, ws.max_row + 1) if _texto(ws.cell(r, 1).value) == _texto(c["id"])]
            if not filas:
                raise CorreccionInvalida(f"no se encontró '{c['id']}' en la columna A de '{c['hoja']}'")
            if len(filas) > 1:
                raise CorreccionInvalida(f"'{c['id']}' aparece {len(filas)} veces en '{c['hoja']}' (filas {filas})")

            celda = ws.cell(filas[0], col)
            actual = _texto(celda.value)
            if c["valor_anterior"] != "*" and actual != _texto(c["valor_anterior"]):
                raise CorreccionInvalida(
                    f"{c['hoja']}!{celda.coordinate} ({c['id']}): se esperaba '{c['valor_anterior']}' y hay '{actual}'"
                )
            nuevo_txt = c["valor_nuevo"].replace("{actual}", "" if actual == VACIO else str(celda.value))
            nuevo = _valor_a_escribir(c["columna"], nuevo_txt)
            cambios.append((celda, actual, nuevo, c))
        except CorreccionInvalida as e:
            errores.append(f"línea {n} del CSV: {e}")

    if errores:
        raise CorreccionInvalida("No se guardó nada:\n  " + "\n  ".join(errores))

    registro = []
    for celda, actual, nuevo, c in cambios:
        celda.value = nuevo
        celda.fill = AMARILLO
        registro.append((c["hoja"], celda.coordinate, c["id"], actual, _texto(nuevo), c["motivo"]))

    hoja_cambios = wb.create_sheet(f"CAMBIOS_{hoy:%d-%m-%Y}", 0)
    hoja_cambios.append(["hoja", "celda", "id / legajo", "valor anterior", "valor nuevo", "motivo"])
    for x in hoja_cambios[1]:
        x.font = Font(bold=True)
    for fila in registro:
        hoja_cambios.append(list(fila))
    for letra, ancho in zip("ABCDEF", (16, 8, 18, 40, 70, 60)):
        hoja_cambios.column_dimensions[letra].width = ancho

    wb.save(ruta_salida)
    return registro


def main():
    ap = argparse.ArgumentParser(description="Aplica correcciones a una copia de un exporte del Sheet.")
    ap.add_argument("excel", help="Exporte .xlsx del Google Sheet (no se modifica)")
    ap.add_argument("correcciones", help="CSV con hoja,id,columna,valor_anterior,valor_nuevo,motivo")
    ap.add_argument("--salida", help="Archivo de salida (default: '<excel> (corregido).xlsx')")
    args = ap.parse_args()

    base, ext = os.path.splitext(args.excel)
    salida = args.salida or f"{base} (corregido){ext}"
    if os.path.abspath(salida) == os.path.abspath(args.excel):
        raise SystemExit("La salida no puede ser el mismo archivo que el exporte original.")
    try:
        registro = aplicar(args.excel, leer_correcciones(args.correcciones), salida)
    except CorreccionInvalida as e:
        raise SystemExit(str(e))
    for fila in registro:
        print(" | ".join(str(x) for x in fila))
    print(f"\n{len(registro)} celdas cambiadas → {salida}")


if __name__ == "__main__":
    main()
